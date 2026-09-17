# -*- coding: utf-8 -*-
"""导出节日生命周期时间点 → docs/产品生命周期时间点0825-最新版.xlsx

格式：把"月份+旬"区间转成纯粹日期区间，如 12/21-12/25、8/1-8/31、12/11-1/10（跨年）。
多个区间用"；"分隔；跨年区间在备注列标注。

用法：python -m scripts.export_lifecycle_timing
"""

import asyncio
import argparse
import calendar
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app.database import async_session_factory  # noqa: E402
from app.models.festival_calendar import FestivalCalendar  # noqa: E402
from app.services.festival_lifecycle import parse_stage  # noqa: E402
from app.services.time_axis import _parse_periods  # noqa: E402

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

STAGES = [
    ("启动期", "launch_stage"),
    ("增长期", "growth_stage"),
    ("热卖期", "hot_period"),
    ("成熟期", "mature_stage"),
    ("下降期", "decline_stage"),
]


def fmt_full_date(dt) -> str:
    return dt.strftime("%Y-%m-%d") if dt else ""


def fmt_periods(raw, fallback_start=None, fallback_end=None) -> str:
    """节日日期/上架时间：支持 单天 / 区间 / 多个分段

    优先 festival_periods（可多段）；无分段时回退 单天 或 start~end 区间。
    """
    periods = _parse_periods(raw)
    if periods:
        parts = []
        for p in periods:
            s = fmt_full_date(p.get("start"))
            e = fmt_full_date(p.get("end"))
            label = str(p.get("label") or "").strip()
            seg = s
            if e and e != s:
                seg += f"~{e}"
            if label:
                seg += f"（{label}）"
            parts.append(seg)
        return "；".join(parts)
    if fallback_start:
        s = fmt_full_date(fallback_start)
        if fallback_end and fallback_end > fallback_start:
            s += f"~{fmt_full_date(fallback_end)}"
        return s
    return ""


def fmt_range(t: tuple) -> str:
    """(start_m, start_d, end_m, end_d) → '12/21-12/25' / '3/17'"""
    sm, sd, em, ed = t
    # 结束日按实际月份天数收敛（2月→28、4月→30…），避免出现 2/31 等不存在的日期
    ed = min(ed, calendar.monthrange(2026, em)[1])
    if (sm, sd) == (em, ed):
        return f"{sm}/{sd}"
    return f"{sm}/{sd}-{em}/{ed}"


def fmt_stage(raw) -> str:
    ranges = parse_stage(raw)
    if not ranges:
        return ""
    return "；".join(fmt_range(r) for r in ranges)


async def main(out_path: str):
    async with async_session_factory() as s:
        rows = (await s.execute(
            select(FestivalCalendar).order_by(FestivalCalendar.festival, FestivalCalendar.festival_date)
        )).scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "生命周期时间点"
    header = ["节日", "节日日期", "预计上架开卖时间", "启动期", "增长期", "热卖期", "成熟期", "下降期", "备注"]
    ws.append(header)
    for c, w in zip("ABCDEFGHI", [18, 42, 18, 20, 20, 20, 20, 20, 32]):
        ws.column_dimensions[c].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in rows:
        note = []
        festival_date = fmt_periods(r.festival_periods, r.festival_date, r.festival_end)
        listing_start = fmt_periods(None, r.listing_start, None)
        if r.listing_start and r.festival_date and r.listing_start.year < r.festival_date.year - 1:
            note.append(f"上架时间疑似旧值（{r.listing_start.year}，节日为{r.festival_date.year}）")
        vals = [r.festival, festival_date, listing_start]
        for name, field in STAGES:
            raw = getattr(r, field, None)
            vals.append(fmt_stage(raw))
            for t in parse_stage(raw):
                if t[0] > t[2]:  # 跨年区间（如 12月中旬-1月上旬）
                    note.append(f"{name}({fmt_range(t)})跨年")
        if r.notes:
            note.append(str(r.notes))
        vals.append("；".join(note))
        ws.append(vals)
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # 冻结首行 + 筛选
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"

    out = out_path or os.path.join(BASE_DIR, "docs", "产品生命周期时间点0825-最新版.xlsx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f"已写入: {out} | 记录数: {len(rows)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", help="输出路径（默认 docs/产品生命周期时间点0825-最新版.xlsx）")
    args = parser.parse_args()
    asyncio.run(main(args.out))
