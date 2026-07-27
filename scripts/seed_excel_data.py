"""
种子数据导入脚本
从Excel文件加载 节日日历 和 产品分类工期 数据到数据库。
"""

import asyncio
import os
import re
import sys
import logging
from datetime import datetime

# 确保能找到 app 模块
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook
from sqlalchemy import select

from app.database import async_session_factory, init_db, engine
from app.models.festival_calendar import FestivalCalendar
from app.models.category_leadtime import CategoryLeadtime

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)


def parse_hot_period(period_str: str) -> tuple:
    """解析热卖期字符串如 '11-12月' 返回 (start_month, end_month)"""
    if not period_str:
        return (None, None)
    period_str = str(period_str).strip().replace(" ", "")
    # 匹配 '11-12月'、'1-3月'、'10-1月'、'整年' 等
    match = re.match(r"(\d+)\s*[-–—]\s*(\d+)", period_str)
    if match:
        return (int(match.group(1)), int(match.group(2)))
    if "整年" in period_str or "全年" in period_str:
        return (1, 12)
    # 单个数字如 '12月'
    match = re.match(r"(\d+)", period_str)
    if match:
        m = int(match.group(1))
        return (m, m)
    return (None, None)


async def import_festival_calendar(xlsx_path: str):
    """导入节日时间+热卖时间段"""
    wb = load_workbook(xlsx_path)
    ws = wb.active
    logger.info(f"读取节日日历文件: {xlsx_path}")

    session = async_session_factory()
    try:
        async with session:
            # 清空已有数据后重新导入
            await session.execute(FestivalCalendar.__table__.delete())
            rows = list(ws.iter_rows(min_row=2, values_only=True))  # 跳过表头
            count = 0
            for row in rows:
                festival = row[0]
                if not festival or not str(festival).strip():
                    continue
                festival = str(festival).strip()

                listing_start = row[1]
                festival_date_val = row[2]
                festival_end = row[3]
                hot_period = str(row[4]).strip() if row[4] else ""
                hot_start, hot_end = parse_hot_period(hot_period)

                # festival_date 可能是 datetime 或文本描述
                festival_date_dt = None
                if isinstance(festival_date_val, datetime):
                    festival_date_dt = festival_date_val
                elif isinstance(festival_date_val, str) and str(festival_date_val).strip():
                    festival_date_dt = None  # 文本描述不存入日期字段
                elif festival_date_val is not None:
                    festival_date_dt = None

                festival_end_dt = None
                if isinstance(festival_end, datetime):
                    festival_end_dt = festival_end
                elif isinstance(festival_end, str) and str(festival_end).strip():
                    festival_end_dt = None

                listing_start_dt = None
                if isinstance(listing_start, datetime):
                    listing_start_dt = listing_start
                elif isinstance(listing_start, str) and str(listing_start).strip():
                    listing_start_dt = None

                # 检查是否已存在（仅按节日名查）
                stmt = select(FestivalCalendar).where(FestivalCalendar.festival == festival)
                result = await session.execute(stmt)
                existing = result.scalars().first()

                if existing:
                    existing.hot_period = hot_period
                    existing.hot_start_month = hot_start
                    existing.hot_end_month = hot_end
                    existing.listing_start = listing_start_dt
                    existing.festival_date = festival_date_dt
                    existing.festival_end = festival_end_dt
                else:
                    record = FestivalCalendar(
                        festival=festival,
                        listing_start=listing_start_dt,
                        festival_date=festival_date_dt,
                        festival_end=festival_end_dt,
                        hot_period=hot_period,
                        hot_start_month=hot_start,
                        hot_end_month=hot_end,
                    )
                    session.add(record)
                count += 1

            await session.commit()
            logger.info(f"节日日历导入完成: {count} 条")
    finally:
        await session.close()


async def import_category_leadtime(xlsx_path: str):
    """导入产品分类工期表"""
    wb = load_workbook(xlsx_path)
    ws = wb.active
    logger.info(f"读取产品分类工期文件: {xlsx_path}")

    session = async_session_factory()
    try:
        async with session:
            # 清空已有数据后重新导入
            await session.execute(CategoryLeadtime.__table__.delete())
            rows = list(ws.iter_rows(min_row=2, values_only=True))  # 跳过表头
            count = 0
            current_level1 = None

            for row in rows:
                level1 = row[0]
                level2 = row[1]
                lead_time_str = row[2]

                # 更新当前一级分类
                if level1 and str(level1).strip():
                    current_level1 = str(level1).strip()

                if not level2 or not str(level2).strip():
                    continue

                level2 = str(level2).strip()
                lead_time_str = str(lead_time_str).strip() if lead_time_str else ""

                # 解析工期字符串如 "15-20天"、"7-10天"、"现货"、"30天"
                lead_time_min, lead_time_max = None, None
                notes = None
                if "现货" in lead_time_str:
                    lead_time_min = 0
                    lead_time_max = 0
                    notes = "现货"
                elif lead_time_str:
                    match = re.match(r"(\d+)\s*[-–—]\s*(\d+)", lead_time_str)
                    if match:
                        lead_time_min = int(match.group(1))
                        lead_time_max = int(match.group(2))
                    else:
                        match = re.match(r"(\d+)", lead_time_str)
                        if match:
                            lead_time_min = int(match.group(1))
                            lead_time_max = int(match.group(1))

                    # 提取备注（中文描述部分）
                    extra = re.sub(r"[\d\s\-–—天左右]", "", lead_time_str).strip()
                    if extra:
                        notes = extra

                # 检查是否已存在
                stmt = select(CategoryLeadtime).where(
                    CategoryLeadtime.level1_category == current_level1,
                    CategoryLeadtime.level2_category == level2,
                )
                result = await session.execute(stmt)
                existing = result.scalar_one_or_none()

                if existing:
                    existing.lead_time_min = lead_time_min
                    existing.lead_time_max = lead_time_max
                    existing.notes = notes
                else:
                    record = CategoryLeadtime(
                        level1_category=current_level1,
                        level2_category=level2,
                        lead_time_min=lead_time_min,
                        lead_time_max=lead_time_max,
                        notes=notes,
                    )
                    session.add(record)
                count += 1

            await session.commit()
            logger.info(f"产品分类工期导入完成: {count} 条")
    finally:
        await session.close()


async def main():
    logger.info("开始导入种子数据...")
    # 先建表
    await init_db()

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    # 导入节日日历
    await import_festival_calendar(os.path.join(base_dir, "docs", "节日时间+热卖时间段.xlsx"))

    # 导入产品分类工期
    await import_category_leadtime(os.path.join(base_dir, "docs", "产品分类工期表.xlsx"))

    logger.info("种子数据导入完成!")
    await engine.dispose()


if __name__ == "__main__":
    asyncio.run(main())
